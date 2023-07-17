from docx import Document
import pandas as pd
import json

# Load the JSON data
with open("data.json") as f:
    data = json.load(f)

# Convert the JSON data to a pandas DataFrame
df = pd.DataFrame(data)

# Create a new Word document
doc = Document()

# Add a table to the Word document
table = doc.add_table(rows=1, cols=len(df.columns))

# Set the headers of the table
for i, column in enumerate(df.columns):
    table.cell(0, i).text = column

# Add the data to the table
for index, row in df.iterrows():
    cells = table.add_row().cells
    for i, value in enumerate(row):
        cells[i].text = str(value)

# Save the Word document
doc.save("data.docx")
